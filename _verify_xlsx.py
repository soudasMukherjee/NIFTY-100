 """
Quick verification script for peer_comparison.xlsx
"""
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent
XLSX_PATH = ROOT / "output" / "peer_comparison.xlsx"

wb = openpyxl.load_workbook(str(XLSX_PATH))
print("Sheets:", len(wb.sheetnames))
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n=== {name} ===")
    print(f"  Dimensions: {ws.max_row} rows x {ws.max_column} cols")
    
    # Header row
    h1 = ws.cell(1, 1).value
    h2 = ws.cell(1, 2).value
    h3 = ws.cell(1, 3).value
    h4 = ws.cell(1, 4).value
    print(f"  Headers: {h1} | {h2} | {h3} | {h4}")
    
    # Check each data row
    benchmark_found = False
    summary_found = False
    for r in range(2, ws.max_row + 1):
        cid = ws.cell(r, 1).value
        fill_rgb = ws.cell(r, 1).fill.start_color.rgb if ws.cell(r, 1).fill.start_color else "none"
        
        if cid and "Peer Group Median" in str(cid):
            summary_found = True
            print(f"  ROW {r}: [SUMMARY] fill={fill_rgb}")
        elif fill_rgb and "F4B183" in str(fill_rgb):
            benchmark_found = True
            print(f"  ROW {r}: [BENCHMARK] company={cid}, fill={fill_rgb}")
        
        # Check percentile color coding on last metric's pctl column
        if ws.max_column >= 4:
            pctl_cell = ws.cell(r, 4)  # ROE (Pctl) column
            pctl_fill = pctl_cell.fill.start_color.rgb if pctl_cell.fill.start_color else "none"
            pctl_val = pctl_cell.value
            if pctl_val is not None and r < ws.max_row:
                print(f"  ROW {r}: company={cid}, ROE Pctl={pctl_val}, fill={pctl_fill}")
    
    print(f"  Benchmark row found: {benchmark_found}")
    print(f"  Summary row found: {summary_found}")

wb.close()
print("\n✅ Verification complete")

