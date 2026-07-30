"""
Sprint 3 · Day 16 — Screener XLSX Exporter

Generates output/screener_output.xlsx with one sheet per preset (6 presets),
20 KPI columns per sheet, sorted by composite_quality_score descending.
Colour-codes cells: green fill for cells meeting the preset threshold,
red fill for cells failing the threshold.

Usage:
    python -m src.screener.screener_exporter
"""

from __future__ import annotations

import logging
import sqlite3
from pathlib import Path
from typing import Any, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.screener.engine import ScreenerEngine, ScreenerRequest

# Setup logging
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
)
logger = logging.getLogger(__name__)

ROOT = Path(__file__).resolve().parents[2]
DB_PATH = ROOT / "nifty100.db"
CONFIG_PATH = ROOT / "src" / "screener" / "screener_config.yaml"
OUTPUT_PATH = ROOT / "output" / "screener_output.xlsx"

# ============================================================
# Colour definitions
# ============================================================
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

# ============================================================
# 20 KPI Columns (as they appear in the merged DataFrame)
# ============================================================
KPI_COLUMNS: list[str] = [
    "company_id",                # 1. Company identifier
    "broad_sector",              # 2. Sector classification
    "composite_quality_score",   # 3. Overall quality score (0-100)
    "return_on_equity_pct",      # 4. ROE (%)
    "debt_to_equity",            # 5. Debt-to-Equity ratio
    "interest_coverage",         # 6. Interest Coverage Ratio
    "free_cash_flow_cr",         # 7. Free Cash Flow (₹Cr)
    "revenue_cagr_5yr",          # 8. Revenue CAGR 5Y (%)
    "pat_cagr_5yr",              # 9. PAT CAGR 5Y (%)
    "net_profit_margin_pct",     # 10. Net Profit Margin (%)
    "operating_profit_margin_pct",  # 11. Operating Profit Margin (%)
    "pe_ratio",                  # 12. P/E Ratio
    "pb_ratio",                  # 13. P/B Ratio
    "dividend_yield_pct",        # 14. Dividend Yield (%)
    "market_cap_crore",          # 15. Market Cap (₹Cr)
    "sales",                     # 16. Sales / Revenue (₹Cr)
    "net_profit",                # 17. Net Profit (₹Cr)
    "earnings_per_share",        # 18. EPS (₹)
    "revenue_cagr_3yr",          # 19. Revenue CAGR 3Y (%)
    "eps_cagr_5yr",              # 20. EPS CAGR 5Y (%)
]

# Column display labels (human-readable)
KPI_LABELS: dict[str, str] = {
    "company_id": "Company",
    "broad_sector": "Sector",
    "composite_quality_score": "Quality Score",
    "return_on_equity_pct": "ROE (%)",
    "debt_to_equity": "D/E",
    "interest_coverage": "ICR",
    "free_cash_flow_cr": "FCF (₹Cr)",
    "revenue_cagr_5yr": "Rev CAGR 5Y (%)",
    "pat_cagr_5yr": "PAT CAGR 5Y (%)",
    "net_profit_margin_pct": "NPM (%)",
    "operating_profit_margin_pct": "OPM (%)",
    "pe_ratio": "P/E",
    "pb_ratio": "P/B",
    "dividend_yield_pct": "Div Yield (%)",
    "market_cap_crore": "Mkt Cap (₹Cr)",
    "sales": "Sales (₹Cr)",
    "net_profit": "Net Profit (₹Cr)",
    "earnings_per_share": "EPS (₹)",
    "revenue_cagr_3yr": "Rev CAGR 3Y (%)",
    "eps_cagr_5yr": "EPS CAGR 5Y (%)",
}

# ============================================================
# Metric key -> DataFrame column mapping
# ============================================================
METRIC_KEY_TO_COLUMN: dict[str, str] = {
    "roe_min": "return_on_equity_pct",
    "de_max": "debt_to_equity",
    "fcf_min": "free_cash_flow_cr",
    "revenue_cagr_5yr_min": "revenue_cagr_5yr",
    "pat_cagr_5yr_min": "pat_cagr_5yr",
    "opm_min": "operating_profit_margin_pct",
    "pe_max": "pe_ratio",
    "pb_max": "pb_ratio",
    "dividend_yield_min": "dividend_yield_pct",
    "icr_min": "interest_coverage",
    "market_cap_min": "market_cap_crore",
    "net_profit_min": "net_profit",
    "eps_cagr_min": "eps_cagr_5yr",
    "asset_turnover_min": "asset_turnover",
    "dividend_payout_max": "dividend_payout_ratio_pct",
    "revenue_cagr_3yr_min": "revenue_cagr_3yr",
    "de_near_zero": "debt_to_equity",
    "sales_min": "sales",
    "de_yoy_change_pct_min": "de_yoy_change_pct",
}

# Metric keys that have special handling
SPECIAL_HANDLING: dict[str, str] = {
    "de_max": "financials_sector_skip",
    "de_near_zero": "financials_sector_skip",  # Also skip Financials for near-zero D/E
    "icr_min": "debt_free_as_infinity",
}

# Direction for each metric key (higher_is_better or lower_is_better)
# Derived from screener_config.yaml
METRIC_DIRECTION: dict[str, str] = {
    "roe_min": "higher_is_better",
    "de_max": "lower_is_better",
    "fcf_min": "higher_is_better",
    "revenue_cagr_5yr_min": "higher_is_better",
    "pat_cagr_5yr_min": "higher_is_better",
    "opm_min": "higher_is_better",
    "pe_max": "lower_is_better",
    "pb_max": "lower_is_better",
    "dividend_yield_min": "higher_is_better",
    "icr_min": "higher_is_better",
    "market_cap_min": "higher_is_better",
    "net_profit_min": "higher_is_better",
    "eps_cagr_min": "higher_is_better",
    "asset_turnover_min": "higher_is_better",
    "dividend_payout_max": "lower_is_better",
    "revenue_cagr_3yr_min": "higher_is_better",
    "de_near_zero": "lower_is_better",
    "sales_min": "higher_is_better",
    "de_yoy_change_pct_min": "higher_is_better",
}


# ============================================================
# Fetch all data (full universe, not filtered)
# ============================================================
def fetch_full_data(conn: sqlite3.Connection) -> pd.DataFrame:
    """
    Load the full Nifty 100 data: financial_ratios, market_cap, P&L, and sectors
    for the latest year per company.
    """
    # Financial ratios + sectors
    ratios_query = """
        SELECT
            fr.company_id,
            fr.year,
            fr.net_profit_margin_pct,
            fr.operating_profit_margin_pct,
            fr.return_on_equity_pct,
            fr.debt_to_equity,
            fr.interest_coverage,
            fr.asset_turnover,
            fr.free_cash_flow_cr,
            fr.fcf_conversion_rate_pct,
            fr.cash_from_operations_cr,
            fr.capex_cr,
            fr.earnings_per_share,
            fr.book_value_per_share,
            fr.dividend_payout_ratio_pct,
            fr.total_debt_cr,
            fr.revenue_cagr_3yr,
            fr.de_yoy_change_pct,
            fr.revenue_cagr_5yr,
            fr.pat_cagr_5yr,
            fr.eps_cagr_5yr,
            fr.composite_quality_score,
            s.broad_sector
        FROM financial_ratios fr
        LEFT JOIN sectors s ON fr.company_id = s.company_id
        WHERE CAST(fr.year AS INTEGER) = (
            SELECT MAX(CAST(fr2.year AS INTEGER)) FROM financial_ratios fr2
            WHERE fr2.company_id = fr.company_id
        )
        ORDER BY fr.company_id, fr.year DESC
    """
    df = pd.read_sql_query(ratios_query, conn)
    df["year"] = df["year"].astype("int64")

    # Market cap
    mcap_query = """
        SELECT
            company_id,
            year,
            market_cap_crore,
            pe_ratio,
            pb_ratio,
            dividend_yield_pct,
            ev_ebitda,
            enterprise_value_crore
        FROM market_cap
        WHERE year = (
            SELECT MAX(year) FROM market_cap mc2
            WHERE mc2.company_id = market_cap.company_id
        )
        ORDER BY company_id, year DESC
    """
    mcap_df = pd.read_sql_query(mcap_query, conn)
    mcap_df["year"] = mcap_df["year"].astype("int64")

    # P&L
    pl_query = """
        SELECT
            company_id,
            year,
            sales,
            net_profit,
            eps,
            operating_profit,
            opm_percentage
        FROM profitandloss
        WHERE period_type = 'ANNUAL'
        AND year = (
            SELECT MAX(year) FROM profitandloss pl2
            WHERE pl2.company_id = profitandloss.company_id
            AND pl2.period_type = 'ANNUAL'
        )
        ORDER BY company_id, year DESC
    """
    pl_df = pd.read_sql_query(pl_query, conn)
    pl_df["year"] = pl_df["year"].astype("int64")

    # Merge
    df = df.merge(mcap_df, on=["company_id", "year"], how="left")
    df = df.merge(
        pl_df[["company_id", "year", "sales", "net_profit", "eps"]],
        on=["company_id", "year"],
        how="left",
        suffixes=("", "_pl"),
    )
    # Use eps from P&L if financial_ratios doesn't have it
    if "earnings_per_share" in df.columns and "eps" in df.columns:
        df["earnings_per_share"] = df["earnings_per_share"].fillna(df["eps"])
        df.drop(columns=["eps"], inplace=True)

    return df


# ============================================================
# Run a single preset filter (using the engine)
# ============================================================
def screen_preset(
    screener: ScreenerEngine,
    preset_name: str,
) -> pd.DataFrame:
    """
    Run the screener for a given preset and return the filtered DataFrame
    sorted by composite_quality_score descending.
    """
    filters = screener.get_preset_filters(preset_name)
    request = ScreenerRequest(filters=filters, preset_name=preset_name)
    result = screener.screen(request)
    return result


# ============================================================
# Apply conditional colour-coding to a worksheet
# ============================================================
def apply_colour_coding(
    ws: Any,
    df: pd.DataFrame,
    preset_filters: dict[str, float],
    col_idx_map: dict[str, int],
    broad_sector_col: int,
) -> None:
    """
    Colour-code cells in the worksheet based on whether the cell value
    meets the preset threshold for that metric.

    Green fill  = value meets the threshold (pass)
    Red fill    = value fails the threshold

    Special cases:
      - financials_sector_skip: skip colouring for Financials sector (D/E)
      - debt_free_as_infinity: NaN values are treated as passing (green)
    """
    for row_idx in range(2, ws.max_row + 1):  # row 1 is header
        company_sector = ws.cell(row=row_idx, column=broad_sector_col).value or ""

        for metric_key, threshold in preset_filters.items():
            if metric_key not in METRIC_KEY_TO_COLUMN:
                continue

            col_name = METRIC_KEY_TO_COLUMN[metric_key]
            if col_name not in col_idx_map:
                continue

            excel_col = col_idx_map[col_name]
            cell = ws.cell(row=row_idx, column=excel_col)
            cell_value = cell.value

            # Skip non-numeric values
            if cell_value is None:
                # For debt_free_as_infinity (ICR), NaN passes
                special = SPECIAL_HANDLING.get(metric_key)
                if special == "debt_free_as_infinity":
                    cell.fill = GREEN_FILL
                continue

            if not isinstance(cell_value, (int, float)):
                continue

            # Handle special cases
            special = SPECIAL_HANDLING.get(metric_key)

            # financials_sector_skip: skip colouring for Financials sector
            if special == "financials_sector_skip" and company_sector == "Financials":
                continue

            # Determine pass/fail based on direction
            direction = METRIC_DIRECTION.get(metric_key, "higher_is_better")

            if direction == "higher_is_better":
                passes = cell_value >= threshold
            else:  # lower_is_better
                passes = cell_value <= threshold

            cell.fill = GREEN_FILL if passes else RED_FILL


# ============================================================
# Write a DataFrame to a worksheet with formatting
# ============================================================
def write_sheet(
    wb: Workbook,
    sheet_name: str,
    df: pd.DataFrame,
    preset_filters: dict[str, float],
    kpi_display_name: str = "",
) -> None:
    """
    Create a worksheet with 20 KPI columns, sorted by composite_quality_score desc,
    with colour-coded cells indicating pass/fail against preset thresholds.
    """
    ws = wb.create_sheet(title=sheet_name)

    # Ensure the KPI columns exist in the DataFrame, fill missing with None
    for col in KPI_COLUMNS:
        if col not in df.columns:
            df[col] = None

    # Sort by composite_quality_score descending
    if "composite_quality_score" in df.columns:
        df = df.sort_values(
            "composite_quality_score", ascending=False, na_position="last"
        ).reset_index(drop=True)

    # Select only the 20 KPI columns
    out_df = df[KPI_COLUMNS].copy()

    # Build column index map: column_name -> excel_column_number (1-based)
    col_idx_map: dict[str, int] = {}
    for i, col in enumerate(KPI_COLUMNS, start=1):
        col_idx_map[col] = i

    # Write header row
    for col_idx, col_name in enumerate(KPI_COLUMNS, start=1):
        label = KPI_LABELS.get(col_name, col_name)
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    # Write data rows
    for row_idx, (_, row) in enumerate(out_df.iterrows(), start=2):
        for col_idx, col_name in enumerate(KPI_COLUMNS, start=1):
            value = row.get(col_name)
            # Format numbers nicely
            if isinstance(value, float):
                # Format percentages (0-100 range) with 1 decimal
                if col_name.endswith("_pct") or col_name == "composite_quality_score":
                    cell_value = round(value, 1)
                else:
                    cell_value = round(value, 2)
            else:
                cell_value = value

            cell = ws.cell(row=row_idx, column=col_idx, value=cell_value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="right" if col_idx > 2 else "left")

            # Special formatting for company_id (bold)
            if col_name == "company_id":
                cell.font = Font(bold=True)

    # Apply colour-coding based on preset filters
    broad_sector_col = col_idx_map.get("broad_sector", 2)
    apply_colour_coding(ws, out_df, preset_filters, col_idx_map, broad_sector_col)

    # Auto-fit column widths (approximate)
    for col_idx, col_name in enumerate(KPI_COLUMNS, start=1):
        label = KPI_LABELS.get(col_name, col_name)
        # Estimate width based on header and data
        max_len = len(str(label))
        # Sample first few data rows for width estimation
        for row_idx in range(2, min(ws.max_row + 1, 6)):
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val is not None:
                val_str = f"{cell_val:.1f}" if isinstance(cell_val, float) else str(cell_val)
                max_len = max(max_len, len(val_str))
        # Cap at reasonable widths
        col_width = min(max_len + 3, 22)
        col_width = max(col_width, 10)
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    # Freeze top row
    ws.freeze_panes = "A2"

    logger.info(
        f"  Sheet '{sheet_name}': {len(out_df)} companies, {len(KPI_COLUMNS)} KPI columns"
    )


# ============================================================
# Main export function
# ============================================================
def export_screener_to_xlsx(
    db_path: Path = DB_PATH,
    output_path: Path = OUTPUT_PATH,
) -> Path:
    """
    Generate output/screener_output.xlsx with one sheet per preset.

    Each sheet contains the companies that pass that preset's filters,
    with 20 KPI columns, sorted by composite_quality_score descending,
    and colour-coded cells.

    Returns the path to the generated file.
    """
    logger.info("=" * 60)
    logger.info("SCREENER XLSX EXPORTER")
    logger.info("=" * 60)

    output_path.parent.mkdir(exist_ok=True)

    # Open database connection for full data fetch
    conn = sqlite3.connect(str(db_path))
    screener: ScreenerEngine | None = None

    try:
        # Fetch full data universe
        logger.info("Loading full Nifty 100 data...")
        full_df = fetch_full_data(conn)
        logger.info(f"  Loaded {len(full_df)} companies with {len(full_df.columns)} columns")

        # Load screener engine
        screener = ScreenerEngine(db_path=db_path)
        screener._get_connection()  # init connection

        # Create Excel workbook
        wb = Workbook()
        # Remove default sheet (if it exists)
        default_sheet = wb.active
        if default_sheet is not None:
            wb.remove(default_sheet)

        # Process each preset
        preset_list = screener.list_presets()
        logger.info(f"\nProcessing {len(preset_list)} presets...")

        for preset_key, preset_name in preset_list:
            logger.info(f"\n--- Preset: {preset_name} ({preset_key}) ---")

            # Get preset filters
            preset_filters = screener.get_preset_filters(preset_key)
            logger.info(f"  Filters: {preset_filters}")

            # Run the screener to get filtered company IDs
            request = ScreenerRequest(filters=preset_filters, preset_name=preset_key)
            filtered_result = screener.screen(request)

            if len(filtered_result) == 0:
                logger.warning(f"  WARNING: Preset '{preset_key}' returned 0 companies!")
                # Still create an empty sheet with headers
                empty_df = pd.DataFrame(columns=KPI_COLUMNS)
                write_sheet(wb, preset_key[:31], empty_df, preset_filters)
                continue

            # Subset the full data to only include companies that passed the preset
            passed_ids = set(filtered_result["company_id"].tolist())
            preset_df = full_df[full_df["company_id"].isin(passed_ids)].copy()

            logger.info(f"  Companies passed: {len(preset_df)}")

            # Write sheet (Excel sheet names max 31 chars)
            sheet_name = preset_key[:31]
            write_sheet(wb, sheet_name, preset_df, preset_filters)

        # Save workbook
        logger.info(f"\nSaving to {output_path}...")
        wb.save(str(output_path))
        logger.info(f"✓ Successfully exported to {output_path}")

        # Print summary
        print(f"\n{'=' * 60}")
        print(f"  SCREENER OUTPUT GENERATED")
        print(f"{'=' * 60}")
        print(f"  File: {output_path}")
        print(f"  Sheets: {len(wb.sheetnames)}")
        for name in wb.sheetnames:
            ws = wb[name]
            print(f"    • {name}: {ws.max_row - 1} companies × {ws.max_column} KPIs")
        print(f"{'=' * 60}\n")

        return output_path

    finally:
        conn.close()
        if screener is not None:
            screener.close()


# ============================================================
# CLI Entry Point
# ============================================================
if __name__ == "__main__":
    export_screener_to_xlsx()

